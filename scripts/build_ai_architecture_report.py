from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\GradProject\IZEE")
OUT = ROOT / "IZEE_AI_ETA_Alert_Engine_Report.docx"


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
    doc.styles[style_name].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Heading 1"].font.size = Pt(16)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
doc.styles["Heading 2"].font.size = Pt(13)
doc.styles["Heading 2"].font.bold = True
doc.styles["Heading 2"].font.color.rgb = RGBColor(47, 84, 150)

if "CodeBlock" not in doc.styles:
    code_style = doc.styles.add_style("CodeBlock", 1)
else:
    code_style = doc.styles["CodeBlock"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
code_style.font.size = Pt(8.5)


def shade_paragraph(paragraph, fill="F3F6FA"):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def code_block(text):
    p = doc.add_paragraph(style="CodeBlock")
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text.strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.5)
    shade_paragraph(p)


def callout(title, body):
    p = doc.add_paragraph()
    shade_paragraph(p, "EAF2F8")
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    p.add_run(body)


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(headers, rows):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        cell._tc.get_or_add_tcPr().append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.8)
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Using AI in IZEE ETA and Alert Engine")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(31, 78, 121)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Recommended AI architecture for prediction, alerting, and explainability").italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

doc.add_heading("1. Short Answer", level=1)
doc.add_paragraph(
    "Yes, IZEE can use AI in both the ETA Engine and the Alert Engine. The important design decision is that AI should play different roles in each component."
)
callout(
    "Recommended principle",
    "ETA can be AI-led because it is a prediction problem. Alert Engine should be rule-led and AI-enhanced because operational alerts must be explainable, consistent, testable, and easy to defend."
)

doc.add_heading("2. Why AI Fits ETA Very Well", level=1)
doc.add_paragraph(
    "ETA is naturally a machine learning problem. The system needs to estimate future arrival time or future delay from historical and live operational features."
)
bullets([
    "ETA asks: when will this vehicle arrive?",
    "ETA can learn from historical segment travel times, speed, time period, route, direction, and current progress.",
    "The model output can be numeric: predicted travel time, predicted arrival timestamp, or predicted delay in seconds.",
    "Because output is measurable, the model can be evaluated with MAE, RMSE, MAPE, and error distribution.",
])

doc.add_heading("3. Why Alert Engine Should Stay Hybrid", level=1)
doc.add_paragraph(
    "The Alert Engine creates operational decisions. These decisions may be shown to passengers, drivers, supervisors, or a control center. Therefore, the core decision must be deterministic and explainable."
)
bullets([
    "A rule like travel_time > 2x average is easy to explain and reproduce.",
    "A rule like 3 delayed vehicles on the same route means disruption is easy to defend in a graduation discussion.",
    "A fully AI-based alert decision needs training data, labels, evaluation metrics, false-positive analysis, and model governance.",
    "AI is still valuable in the Alert Engine, but mainly for enhancement: explanation, prioritization, anomaly score, probable cause, and recommended action.",
])

doc.add_heading("4. Recommended Architecture", level=1)
code_block("""
Vehicle Observations
        ↓
Vehicle State Engine
        ↓
Event Engine
        ↓
Segment Statistics
        ↓
ETA Engine AI/ML
        ↓
Rule-Based Alert Engine
        ↓
AI Enhancement Layer
        ↓
alerts table / APIs / Control Center
""")

doc.add_heading("5. AI Usage by Component", level=1)
table(
    ["Component", "Best AI Role", "Recommended Model Type", "Output"],
    [
        ["ETA Engine", "Main prediction model", "XGBoost, LightGBM, Random Forest, LSTM/GRU later", "predicted_arrival_time, predicted_travel_time, predicted_delay_seconds"],
        ["Alert Engine", "Decision support/enhancement", "Rules + anomaly model + optional LLM", "alert type, severity, explanation, recommendation, confidence"],
        ["AI Assistant Layer", "Natural language generation", "LLM such as GPT, Claude, Gemini, or local Llama", "human-readable summaries and supervisor recommendations"],
    ],
)

doc.add_heading("6. Recommended Models", level=1)
table(
    ["Use Case", "Model", "Why It Fits", "Complexity"],
    [
        ["ETA prediction", "XGBoost Regressor", "Strong for tabular features, fast, explainable, good for graduation project", "Medium"],
        ["ETA prediction", "LightGBM", "Similar to XGBoost and efficient on large datasets", "Medium"],
        ["ETA prediction", "Random Forest Regressor", "Simple baseline and easy to explain", "Low"],
        ["Anomaly detection", "Isolation Forest", "Detects unusual segment travel times without heavy labels", "Low/Medium"],
        ["Alert text explanation", "LLM", "Turns technical alert data into clear passenger/control-center messages", "Low if API/local LLM exists"],
        ["Advanced sequence ETA", "LSTM / GRU / Temporal Fusion Transformer", "Learns time-series patterns, but needs more data and ML work", "High"],
    ],
)

doc.add_heading("7. ETA Engine Feature Design", level=1)
doc.add_paragraph("Suggested model inputs:")
bullets([
    "route_id",
    "segment_id",
    "direction",
    "day_of_week / day_type",
    "time_period",
    "current_speed",
    "segment_progress",
    "distance_to_next_stop",
    "avg_travel_time",
    "std_travel_time",
    "sample_count",
    "previous_segment_delay",
    "current_delay",
])
doc.add_paragraph("Suggested model outputs:")
bullets([
    "predicted_travel_time",
    "predicted_arrival_time",
    "predicted_delay_seconds",
    "confidence_score",
])

doc.add_heading("8. Alert Engine AI Enhancement Examples", level=1)
doc.add_paragraph("The rule-based Alert Engine creates the alert first:")
code_block("""
Input:
travel_time = 360 seconds
avg_travel_time = 120 seconds
route_id = CTA_M_112
segment_id = 679_1994

Rule:
360 / 120 = 3.0x normal travel time
→ high delay alert
""")
doc.add_paragraph("Then AI can enhance it:")
code_block("""
AI summary:
Severe delay detected on CTA_M_112 between stops 679 and 1994.
The vehicle is taking about 6 minutes instead of the usual 2 minutes.

AI probable cause:
Possible congestion or operational blockage on this segment.

AI recommendation:
Notify passengers and monitor nearby vehicles on the same route.
""")

doc.add_heading("9. Why Not Fully AI-Based Alerts?", level=1)
table(
    ["Concern", "Why It Matters", "Hybrid Solution"],
    [
        ["Explainability", "Supervisors and examiners can ask why an alert was created.", "Rules provide the exact condition; AI explains it."],
        ["Consistency", "The same input should produce the same operational decision.", "Rules stay deterministic."],
        ["Testing", "Rule thresholds can be unit-tested easily.", "AI output is evaluated separately."],
        ["Safety", "Wrong alerts can confuse passengers or supervisors.", "AI does not replace the core decision."],
        ["Viva defense", "A rule-based core is easier to defend academically.", "AI is positioned as enhancement and prediction."],
    ],
)

doc.add_heading("10. Final Recommended Design", level=1)
callout(
    "Final recommendation",
    "Use AI heavily in ETA prediction, and use AI carefully in the Alert Engine as an enhancement layer. This gives the system intelligence without losing explainability."
)
code_block("""
ETA Engine:
AI-led prediction
    → predicted delay / arrival time

Alert Engine:
Rule-led decision
    → delay / dwell_issue / disruption

AI Enhancement:
LLM or ML support
    → explanation / priority / probable cause / recommendation
""")

doc.add_heading("11. Suggested Implementation Roadmap", level=1)
table(
    ["Phase", "Task", "Result"],
    [
        ["Phase 1", "Keep current rule-based Alert Engine stable.", "Reliable deterministic alerts."],
        ["Phase 2", "Build XGBoost ETA baseline.", "Predicted arrival time and predicted delay."],
        ["Phase 3", "Feed predicted_delay into Alert Engine.", "Early warning delay alerts."],
        ["Phase 4", "Add LLM-based alert explanation.", "Cleaner messages for passengers/supervisors."],
        ["Phase 5", "Add anomaly/confidence scoring.", "Better prioritization and alert trust level."],
    ],
)

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.text = "IZEE AI Architecture Recommendation"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)

doc.save(OUT)
print(OUT)
