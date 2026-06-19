from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\GradProject\IZEE")
OUT = ROOT / "IZEE_Technical_Implementation_Report.docx"

FILES_FOR_REFS = {
    "Event detectors": "event_engine/detectors.py",
    "Event engine orchestrator": "event_engine/engine.py",
    "Traversal lifecycle": "event_engine/traversal_lifecycle.py",
    "Dwell lifecycle": "event_engine/dwell_lifecycle.py",
    "Observation pipeline": "services/observation_pipeline.py",
    "Bulk replay": "scripts/bulk_replay.py",
    "GTFS observation generator": "scripts/generate_gtfs_pipeline_observations.py",
    "Schema sync": "scripts/sync_db_schema.py",
    "GTFS pipeline CMD runner": "scripts/run_gtfs_pipeline_test.cmd",
    "Alert engine": "alert_engine/engine.py",
    "Alert rules": "alert_engine/rules.py",
    "Alert aggregator": "alert_engine/aggregator.py",
    "Alert deduplicator": "alert_engine/deduplicator.py",
}


def read(rel):
    path = ROOT / rel
    return path.read_text(encoding="utf-8", errors="replace") if path.exists() else ""


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
    doc.styles[style_name].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.2)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].font.size = Pt(16)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
doc.styles["Heading 2"].font.size = Pt(13)
doc.styles["Heading 2"].font.bold = True
doc.styles["Heading 2"].font.color.rgb = RGBColor(47, 84, 150)
doc.styles["Heading 3"].font.size = Pt(11.5)
doc.styles["Heading 3"].font.bold = True
doc.styles["Heading 3"].font.color.rgb = RGBColor(68, 68, 68)

if "CodeBlock" not in doc.styles:
    code_style = doc.styles.add_style("CodeBlock", 1)
else:
    code_style = doc.styles["CodeBlock"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
code_style.font.size = Pt(8)
code_style.paragraph_format.space_before = Pt(2)
code_style.paragraph_format.space_after = Pt(7)


def shade_paragraph(p, fill="F3F6FA"):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shading)


def add_code(text, fill="F3F6FA"):
    p = doc.add_paragraph(style="CodeBlock")
    run = p.add_run(text.strip() if text.strip() else "# empty")
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8)
    shade_paragraph(p, fill)


def add_callout(title, body, fill="EAF2F8"):
    p = doc.add_paragraph()
    shade_paragraph(p, fill)
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    p.add_run(body)


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        tc_pr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(8.3)
    doc.add_paragraph()
    return table


def add_toc():
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)
    doc.add_paragraph("Right-click the table of contents in Word and choose Update Field to refresh page numbers.")


def title_page():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IZEE Technical Implementation Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Event Engine, Service Observation Pipeline, and Alert Engine Integration Handoff")
    r.italic = True
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font.size = Pt(10)
    add_callout(
        "Scope",
        "This document is based on the current D:\\GradProject\\IZEE repository contents, implemented scripts, "
        "database observations, and the latest GTFS-based CMD pipeline test. Git commit references were unavailable "
        "because Git reported a Windows dubious-ownership safety block in this sandbox, so references are given by file path."
    )


title_page()
doc.add_page_break()
doc.add_heading("Table of Contents", level=1)
add_toc()
doc.add_page_break()

doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "IZEE is a transit management pipeline that ingests simulated or vehicle-location observations, matches them to "
    "GTFS route geometry, derives vehicle state, detects operational events, aggregates segment statistics, and generates "
    "passenger/control-center alert decisions."
)
add_callout(
    "Current result",
    "The GTFS-based CMD test runs through generated observations, Vehicle State Engine, Event Engine, segment statistics, "
    "and Alert Engine. The latest run processed 88 observations, created 30 transit events, computed 1 segment-statistics "
    "cell, and inserted 4 alerts: 3 medium delay alerts and 1 medium disruption alert."
)
bullets([
    "Main objective: convert vehicle observations into reliable operational decisions.",
    "Primary database: PostgreSQL database izee_db on localhost:5433 for this machine.",
    "Core pipeline: Observation source -> Service Observation Pipeline -> Vehicle State Engine -> Event Engine -> transit_events -> segment_statistics -> Alert Engine -> alerts.",
    "Current limitation: dwell_issue alert generation is blocked because the full pipeline test still does not produce dwell_time events.",
])

doc.add_heading("2. Current System Architecture", level=1)
doc.add_heading("2.1 High-Level Architecture Diagram", level=2)
add_code("""flowchart TD
    A[SUMO / Generated / API Observations] --> B[Validation and Normalization]
    B --> C[Service Observation Pipeline]
    C --> D[Vehicle State Engine]
    D --> E[Event Engine Detectors]
    E --> F[Traversal and Dwell Lifecycle]
    F --> G[(transit_events)]
    G --> H[segment_stats.py]
    H --> I[(segment_statistics)]
    G --> J[Alert Engine]
    I --> J
    J --> K[(alerts)]
    K --> L[Passenger / Driver / Supervisor / Control APIs]""")
doc.add_heading("2.2 Component Descriptions", level=2)
add_table(["Component", "Files", "Responsibility"], [
    ["Input and replay", "scripts/bulk_replay.py, scripts/generate_gtfs_pipeline_observations.py", "Loads JSONL observations, normalizes flat or nested location formats, and replays them through the service pipeline."],
    ["Validation", "utils/validators.py, schemas/*.py", "Checks timestamps, lat/lon, speed, and bearing before state processing."],
    ["Reference loading", "reference/loader.py, gtfs/*", "Loads route-specific stops and segments using route_id and direction."],
    ["Vehicle State Engine", "vehicle_state/*.py", "Map-matches observations, computes progress, stop ownership, movement state, and live state."],
    ["Event Engine", "event_engine/*.py", "Detects stop arrivals/departures, segment transitions, traversal completion, and dwell lifecycle events."],
    ["Segment statistics", "scripts/segment_stats.py", "Aggregates segment_completed events into avg/median/std/sample_count cells."],
    ["Alert Engine", "alert_engine/*.py", "Turns events and statistics into delay, dwell_issue, and disruption alerts."],
    ["Database", "models/*.py, database/connection.py", "Stores observations, live state, transit events, segment stats, and alerts."],
])
doc.add_heading("2.3 Dependencies and Integrations", level=2)
bullets([
    "Python runtime with SQLAlchemy and psycopg2 connectivity.",
    "PostgreSQL on localhost:5433, database izee_db, user postgres, local password 1234.",
    "GTFS static files under gtfs/: agency, routes, trips, stops, stop_times.",
    "Environment overrides: IZEE_CONVERTED_DIR and IZEE_REPLAY_FILES for replay file selection.",
])

doc.add_heading("3. Work Completed", level=1)
add_table(["Order", "Completed work", "Before", "Change made", "Why required", "Expected impact"], [
    ["1", "PostgreSQL local connection fixed", "Project expected localhost:5432.", "database/connection.py now points at localhost:5433.", "Postgres service accepted connections on 5433.", "Application can connect to local izee_db."],
    ["2", "Alert Engine implemented", "alert_engine files were empty skeletons.", "Added model, config, rules, deduplicator, aggregator, engine, and runner.", "Needed to satisfy the Alert Engine build document.", "Batch alert generation from transit_events and segment_statistics."],
    ["3", "Alert table registration", "init_db.py did not import alert_engine.models.", "Added import alert_engine.models.", "Required Base.metadata.create_all to include alerts.", "alerts table can be created through normal init."],
    ["4", "Smoke test created", "Only manual DB checks existed.", "Added scripts/alert_engine_smoke_test.py.", "Needed isolated Alert Engine verification.", "Verified delay, dwell_issue, disruption, and dedup logic in isolation."],
    ["5", "GTFS-based observation generator", "Smoke test used artificial SMOKE_* database rows.", "Added generator using CTA_M_112 GTFS stops.", "Needed realistic pipeline input matching user observation shape.", "Generated flat lat/lon/speed/bearing rows from GTFS route coordinates."],
    ["6", "Bulk replay input flexibility", "Hardcoded external path and nested speed_kmh/location.", "Added IZEE_CONVERTED_DIR, IZEE_REPLAY_FILES, flat lat/lon/speed support.", "Required CMD testing with local generated data.", "Replay can use generated local JSONL."],
    ["7", "Database schema sync", "Existing DB tables were older than SQLAlchemy models.", "Added scripts/sync_db_schema.py.", "create_all does not alter existing tables.", "Unblocked replay against existing local DB."],
    ["8", "GTFS test runner", "Pipeline steps had to be run manually.", "Added scripts/run_gtfs_pipeline_test.cmd.", "User wanted CMD-visible testing.", "Single command runs generation, init, schema sync, replay, stats, alerts, SQL verification."],
    ["9", "Event departure lineage improvement", "stop_departure used previous next_stop_id only.", "Pipeline carries current_stop_id and detector prefers it for departure.", "Dwell lifecycle needs departure stop to match arrival stop.", "Improves correctness, but dwell_time remains blocked by stop ownership behavior."],
])

doc.add_heading("4. Event Engine Detectors - Changes, Rationale, and Outcomes", level=1)
add_table(["Detector", "Purpose", "Inputs", "Detection logic", "Failure handling", "Change status"], [
    ["detect_stop_events", "Detects stop_arrival and stop_departure from movement_state transitions.", "current_state, previous_state", "arrival: curr=at_stop and prev!=at_stop; departure: prev=at_stop and curr!=at_stop.", "Returns empty list if no transition.", "Changed departure stop_id to prefer previous_state.current_stop_id."],
    ["detect_segment_transition", "Detects forward, skipped, and reverse segment transitions.", "current_state, previous_state, route_reference", "same sequence=no event; small reverse ignored; larger reverse emits low-confidence segment_travel; forward gaps emit segment_travel.", "Missing stop map entries are skipped.", "No direct change; documented as central behavior."],
    ["detect_delay_event", "Detects delay level transitions from current_delay.", "current_state, previous_state", "Maps delay into minor/moderate/severe and emits only when level changes.", "No event if delay values missing.", "No direct change; Alert Engine does not depend on these delay events."],
])
doc.add_heading("4.1 Detector Processing Flow", level=2)
add_code("""sequenceDiagram
    participant VS as Vehicle State
    participant EE as Event Engine
    participant DS as detect_stop_events
    participant DT as detect_segment_transition
    participant DD as detect_delay_event
    participant BL as build_event
    participant TL as traversal_lifecycle
    participant DL as dwell_lifecycle

    VS->>EE: current_state + previous_state + route_reference
    EE->>DS: movement_state transition
    EE->>DT: stop_sequence transition
    EE->>DD: current_delay transition
    DS-->>EE: raw stop events
    DT-->>EE: raw segment_travel events
    DD-->>EE: raw delay events
    EE->>BL: standardize raw events
    EE->>TL: generate segment_completed
    EE->>DL: generate dwell_time when arrival/departure match
    EE-->>DB: TransitEvent rows""")
doc.add_heading("4.2 Original Implementation and Limitations", level=2)
bullets([
    "Stop departure used previous_state.next_stop_id, which can identify the next stop rather than the stop where a vehicle was dwelling.",
    "The Event Engine emits print-heavy debug output rather than structured logs.",
    "detect_delay_event emits event_type delay, but the Alert Engine design relies on segment_completed and segment_statistics.",
    "Dwell lifecycle requires stop_arrival and stop_departure to share the same stop_id; stop ownership changes near segment boundaries can drop dwell_time.",
])
doc.add_heading("4.3 Changes Implemented", level=2)
add_code("""# services/observation_pipeline.py
previous_state = {
    \"progress\": db_state.progress,
    \"movement_state\": db_state.movement_state,
    \"current_stop_id\": db_state.current_stop_id,
    \"next_stop_id\": db_state.next_stop_id,
    \"segment_id\": db_state.segment_id,
    \"stop_sequence\": db_state.stop_sequence,
    \"current_delay\": db_state.current_delay,
}

# event_engine/detectors.py
\"stop_id\": (
    previous_state.get(\"current_stop_id\")
    or previous_state.get(\"next_stop_id\")
)""")
doc.add_heading("4.4 Before vs After Comparison", level=2)
add_table(["Previous behavior", "Updated behavior", "Reason for change", "Result achieved"], [
    ["stop_departure stop_id came from previous next_stop_id only.", "stop_departure prefers previous current_stop_id and falls back to next_stop_id.", "Dwell lifecycle needs departure from the same stop that generated arrival.", "Improved lineage; dwell_time still blocked by stop ownership in current generated scenario."],
    ["Previous state carried limited fields.", "Previous state includes current_stop_id and segment_id.", "Detectors and lifecycle processors need stop/segment ownership context.", "Improves correctness of lifecycle decisions."],
    ["Debug prints were scattered.", "Still present and documented as technical debt.", "Structured logging was outside this change set.", "Operational visibility remains noisy but useful locally."],
])

doc.add_heading("5. Service/Observation Pipeline - Changes, Rationale, and Outcomes", level=1)
add_table(["Stage", "Inputs", "Outputs", "Internal logic", "Failure handling/performance"], [
    ["Ingestion", "API payload or JSONL raw observation.", "Standard observation dict.", "Normalizes timestamp, lat/lon, speed, bearing, route_id, direction, scenario metadata.", "Malformed JSON lines are skipped in bulk replay."],
    ["Validation", "Observation fields.", "Validated fields.", "validate_location, validate_speed, validate_bearing.", "Raises exception and rolls back current batch if processing fails."],
    ["Route reference routing", "route_id + direction.", "route_reference with stops and segments.", "route_cache keyed by (route_id, direction).", "Missing trip raises route-reference exception."],
    ["Previous state loading", "vehicle_id.", "previous_state dict.", "Uses vehicle_state_cache in bulk replay, otherwise DB query.", "Cache avoids per-observation DB round trip."],
    ["Vehicle state processing", "observation + previous_state + route_reference.", "current vehicle state.", "Map matching, projection, progress stabilization, movement-state detection.", "validate_vehicle_state clamps invalid state."],
    ["Event processing", "current state + previous state + route reference.", "TransitEvent dictionaries.", "Runs detectors, traversal lifecycle, dwell lifecycle.", "Duplicate segment_travel filtered by last_transition_per_vehicle."],
    ["Persistence", "observation mapping, events, live state.", "transit_observations, transit_events, vehicle_live_state.", "bulk_insert_mappings for observations; db.add for events/live state.", "Savepoint per batch; failed batch rolled back and counted."],
])
doc.add_heading("5.1 Updated Sequence Diagram", level=2)
add_code("""sequenceDiagram
    participant R as bulk_replay.py
    participant P as observation_pipeline
    participant Ref as reference.loader
    participant VS as vehicle_state.engine
    participant EE as event_engine.engine
    participant DB as PostgreSQL

    R->>R: read JSONL observation
    R->>R: normalize flat lat/lon/speed or nested location/speed_kmh
    R->>P: process_observation_pipeline(observation, db, cache, persist=False)
    P->>Ref: load_route_reference(route_id, direction)
    Ref-->>P: route stops + segments
    P->>P: load previous_state from vehicle_state_cache
    P->>VS: process_observation()
    VS-->>P: current vehicle state
    P->>EE: process_event(current_state, previous_state, route_reference)
    EE-->>P: event list
    P->>DB: add TransitEvent objects
    P->>DB: update VehicleLiveState
    R->>DB: bulk insert TransitObservation mappings
    R->>DB: commit batch""")
doc.add_heading("5.2 Pipeline Changes Analysis", level=2)
add_table(["Change", "Original behavior", "Updated implementation", "Why changed", "Result"], [
    ["Optional DB/session ownership", "Pipeline always created/owned its own DB session.", "process_observation_pipeline accepts db and owns_session flag.", "Bulk replay needs batched transactions.", "Supports API and replay paths."],
    ["Vehicle state cache", "DB lookup per observation.", "vehicle_state_cache optional dict keyed by vehicle_id.", "Improve replay throughput.", "Replay can process batches with less DB overhead."],
    ["persist_observation flag", "Pipeline added observation ORM object per call.", "Bulk replay passes persist_observation=False and bulk-inserts mappings.", "Avoid one db.add per observation.", "Improved replay performance."],
    ["Route cache key", "Older comments indicate route-only cache.", "Current cache key is (route_id, direction).", "Direction-specific trips have different stop sequences.", "Prevents wrong reference on same route opposite direction."],
    ["Previous state enrichment", "Previous_state had progress, movement_state, next_stop_id, stop_sequence, current_delay.", "Added current_stop_id and segment_id.", "Needed by stop departure and lifecycle lineage.", "Improves stop ownership context."],
    ["Bulk replay flat format support", "Expected nested location and speed_kmh.", "Accepts flat lat/lon/speed and fills defaults.", "User observation format is flat.", "Generated GTFS observations run through real pipeline."],
    ["Day reset", "State could leak across replay days.", "reset_day_state clears trackers and live state.", "Vehicle IDs repeat across days.", "Reduces stale state corruption."],
])

doc.add_heading("6. Changes Made to the Pipeline", level=1)
add_table(["Modification", "Original behavior", "New behavior", "Risks addressed", "Benefits gained"], [
    ["IZEE_CONVERTED_DIR", "Hardcoded external E: path.", "Environment override selects local replay directory.", "Missing external dataset path blocked tests.", "Reproducible local CMD testing."],
    ["IZEE_REPLAY_FILES", "Default TEST_FILES expected fixed files.", "Environment variable constrains replay to generated file.", "Missing day_5/day_6 files would fail.", "Single-file integration test."],
    ["Schema sync script", "create_all did not alter older tables.", "ADD COLUMN IF NOT EXISTS for local tables.", "Undefined column failures.", "Existing local DB can run newer code."],
    ["GTFS generator", "Smoke test inserted fake DB rows directly.", "Generated observations from real GTFS CTA_M_112 stops.", "Smoke data bypassed pipeline.", "Tests Vehicle State and Event Engine path."],
    ["Cleanup script", "Repeated tests could mix previous rows.", "Deletes GTFS_* observations/events/live state and relevant alerts/stats.", "Dedup and stale rows skewed results.", "Repeatable integration test."],
])

doc.add_heading("7. Alert Engine Analysis - Why the Alert Engine Is Blocked From Other Components", level=1)
doc.add_paragraph("The Alert Engine is implemented and works when expected input events exist. The integration blockers are upstream data, schema, and event-production issues that prevent the Alert Engine from receiving complete inputs.")
add_table(["Blocker", "Component", "Technical evidence", "Impact", "Severity", "Recommended fix", "Estimated effort"], [
    ["Missing real replay dataset", "bulk_replay.py / observation source", "Original configured path E:\\Last Semster\\IZEE_SUMO\\converted was not found. A local GTFS generator was created as substitute.", "Cannot validate production-scale SUMO replay.", "High", "Locate/regenerate real converted JSONL files or standardize generated fixtures.", "0.5-1 day once files are available."],
    ["Older database schema", "PostgreSQL tables", "transit_observations and vehicle_live_state missed model columns until sync script.", "Replay failed before Event Engine could create events.", "High", "Use migrations or rebuild DB; keep sync_db_schema.py as dev bridge.", "1-2 days for migration baseline."],
    ["dwell_time not produced in full test", "Event Engine dwell lifecycle / stop ownership", "Latest run: dwell events checked = 0; alerts have delay/disruption only.", "Alert Engine cannot create dwell_issue alerts in real pipeline.", "High", "Fix stop ownership around boundaries and add Dwell.json lifecycle tests.", "1-2 days."],
    ["Segment statistics sample threshold below target", "segment_stats.py / generated data", "sample_count=14 for CTA_M_112 weekday peak segment 679_1994; target is 30.", "Stats are integration-grade but below reliability target.", "Medium", "Generate >=30 baseline samples or use real replay data.", "0.5 day."],
    ["DB-heavy Alert Engine lookup", "alert_engine/engine.py", "_fetch_segment_stats queries DB inside loop.", "May be slow for millions of events and couples decisions to storage.", "Medium", "Prefetch segment_statistics cache.", "0.5-1 day."],
    ["No temporal window for disruption", "alert_engine/aggregator.py", "Groups delay alerts by route without event-time clustering.", "Can group delays across long periods.", "Medium", "Group by route plus rolling 10-15 minute window.", "0.5 day."],
])

doc.add_heading("8. End-to-End Event Flow", level=1)
add_code("""flowchart LR
    S[Observation Source] --> V[Validation]
    V --> R[Route Reference]
    R --> VS[Vehicle State]
    VS --> DE[Detectors]
    DE --> TL[Traversal Lifecycle]
    DE --> DL[Dwell Lifecycle]
    TL --> TE[(transit_events)]
    DL --> TE
    TE --> SS[segment_stats]
    SS --> ST[(segment_statistics)]
    TE --> AE[Alert Engine]
    ST --> AE
    AE --> AL[(alerts)]
    AL --> API[Notification/Consumer APIs]

    V -. fail .-> F1[Validation exception]
    R -. fail .-> F2[Missing route reference]
    VS -. fail .-> F3[Bad state / matching]
    DE -. drop .-> F4[No detector condition]
    DL -. drop .-> F5[Arrival/departure mismatch]
    SS -. weak .-> F6[Low sample_count]
    AE -. skip .-> F7[Missing stats / dedup]""")
doc.add_paragraph("Known failure/drop locations: validation errors, route-reference lookup failures, batch rollback in bulk replay, no detector condition matched, traversal/dwell lifecycle mismatch, low segment-stat samples, missing segment_statistics rows, and Alert Engine deduplication.")

doc.add_heading("9. Technical Debt and Risks", level=1)
add_table(["Risk area", "Risk", "Evidence", "Mitigation"], [
    ["Architecture", "Alert engine orchestration mixes SQL loading, rule evaluation, deduplication, and storage.", "alert_engine/engine.py contains fetch, loop, dedup, and insert logic.", "Split into data_loader.py, repository.py, pure rules.py."],
    ["Reliability", "Dwell lifecycle is fragile near stop/segment boundary ownership.", "GTFS test still produced zero dwell_time events.", "Add deterministic dwell unit tests and improve stop ownership tracking."],
    ["Scalability", "Segment stats queried per segment_completed event.", "_fetch_segment_stats called inside loop.", "Prefetch stats cache keyed by segment_id/route/direction/day_type/time_period."],
    ["Data quality", "Sample_count below target in generated integration test.", "segment_stats.py warning: sample_count 14, samples_needed 16.", "Generate more samples or validate with full replay dataset."],
    ["Observability", "Debug prints instead of structured logging.", "Multiple print statements in detectors and pipeline.", "Use logging module with route/vehicle/event context."],
    ["Database management", "No migrations; schema drift required sync script.", "Undefined column errors surfaced during CMD run.", "Introduce Alembic migrations or controlled rebuild scripts."],
    ["Security/config", "DB password and port hardcoded.", "database/connection.py contains local URL.", "Move DATABASE_URL to environment variable."],
])

doc.add_heading("10. Recommendations", level=1)
add_table(["Priority", "Recommendation", "Expected impact", "Complexity", "Estimated effort"], [
    ["Priority 1 - Critical", "Fix dwell_time event generation with tests covering arrival, dwell, departure at the same GTFS stop.", "Unblocks dwell_issue alerts in full pipeline.", "Medium", "1-2 days."],
    ["Priority 1 - Critical", "Create database migrations or a clean rebuild path.", "Removes schema drift failures.", "Medium", "1-2 days."],
    ["Priority 1 - Critical", "Run against real SUMO converted replay files.", "Validates actual system behavior.", "Depends on data availability", "0.5 day after data is available."],
    ["Priority 2 - High", "Refactor Alert Engine into data_loader, repository, pure rules, aggregator, engine.", "Improves maintainability and future streaming readiness.", "Medium", "1 day."],
    ["Priority 2 - High", "Add temporal-window disruption aggregation.", "Prevents false route-level disruptions.", "Low/Medium", "0.5 day."],
    ["Priority 2 - High", "Add confidence scoring based on sample_count/std_travel_time.", "Makes alert reliability explainable.", "Low", "0.5 day."],
    ["Priority 3 - Medium", "Replace debug prints with structured logging.", "Improves operational visibility.", "Low/Medium", "1 day."],
    ["Priority 3 - Medium", "Move DB settings to environment variables.", "Improves portability and security hygiene.", "Low", "0.25 day."],
])

doc.add_heading("11. Appendix", level=1)
doc.add_heading("11.1 Important Files and Folders", level=2)
add_table(["Path", "Purpose"], [[v, k] for k, v in FILES_FOR_REFS.items()] + [
    ["gtfs/", "Static GTFS reference CSV files used for route/stop extraction."],
    ["local_replay_data/day_1_monday_observations.jsonl", "Generated GTFS-based observations used by the CMD integration test."],
    ["models/", "SQLAlchemy table models."],
])
doc.add_heading("11.2 Environment Variables and Commands", level=2)
add_code("""set PYTHONPATH=D:\\GradProject\\IZEE
set IZEE_CONVERTED_DIR=D:\\GradProject\\IZEE\\local_replay_data
set IZEE_REPLAY_FILES=day_1_monday_observations.jsonl
set PGPASSWORD=1234

cd /d D:\\GradProject\\IZEE
scripts\\run_gtfs_pipeline_test.cmd""")
doc.add_heading("11.3 Latest CMD Test Evidence", level=2)
add_code("""Observations processed: 88
transit_events created: 30
segment_completed events created: 14
Errors: 0

Segment statistics:
CTA_M_112 / weekday / peak / 679_1994 -> sample_count 14

Alert Engine:
Segment events checked: 14
Dwell events checked: 0
Alerts created: 4
Delay alerts: 3
Disruptions: 1

Final alert distribution:
delay      medium   3
disruption medium   1""")
doc.add_heading("11.4 Selected Code References", level=2)
for label, rel in FILES_FOR_REFS.items():
    doc.add_heading(label + " - " + rel, level=3)
    text = read(rel)
    if len(text) > 2200:
        text = text[:2200] + "\n\n... [truncated in report; see repository file for full source]"
    add_code(text, fill="F7F7F7")

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.text = "IZEE Technical Implementation Report - Generated handoff document"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(120, 120, 120)

doc.save(OUT)
print(OUT)
print(OUT.exists())
print(OUT.stat().st_size)
